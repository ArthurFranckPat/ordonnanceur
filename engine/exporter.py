from __future__ import annotations

import os
from typing import Dict

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import Outline

FILL_VERT = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
FILL_ORANGE = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
FILL_ROUGE = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
FILL_SE = PatternFill(start_color="D9E2F3", end_color="D9E2F3", fill_type="solid")
FONT_VERT = Font(color="006100", bold=True)
FONT_ORANGE = Font(color="9C5700", bold=True)
FONT_ROUGE = Font(color="9C0006", bold=True)
FONT_SE = Font(color="2F5496", italic=True)
FONT_HEADER = Font(bold=True, color="FFFFFF", size=10)
FILL_HEADER = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
FONT_LINK = Font(color="0563C1", underline="single")
THIN_BORDER = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin"),
)
FEU_STYLES: Dict[str, tuple] = {
    "VERT": (FILL_VERT, FONT_VERT),
    "ORANGE": (FILL_ORANGE, FONT_ORANGE),
    "ROUGE": (FILL_ROUGE, FONT_ROUGE),
}


def _style_feu(ws, row: int, col: int, val: str) -> None:
    if val in FEU_STYLES:
        f, fo = FEU_STYLES[val]
        ws.cell(row=row, column=col).fill = f
        ws.cell(row=row, column=col).font = fo


def _fmt_sheet(ws, nc: int) -> None:
    for col in range(1, nc + 1):
        c = ws.cell(row=1, column=col)
        c.font = FONT_HEADER
        c.fill = FILL_HEADER
        c.alignment = Alignment(horizontal="center", wrap_text=True)
    for col in range(1, nc + 1):
        ml = 0
        for row in range(1, min(ws.max_row + 1, 100)):
            v = ws.cell(row=row, column=col).value
            if v:
                ml = max(ml, len(str(v)))
        ws.column_dimensions[get_column_letter(col)].width = min(ml + 3, 30)
    for row in range(1, ws.max_row + 1):
        for col in range(1, nc + 1):
            ws.cell(row=row, column=col).border = THIN_BORDER


def export_xl(
    df_cmd: pd.DataFrame,
    df_det: pd.DataFrame,
    df_plan: pd.DataFrame,
    df_crit: pd.DataFrame,
    params: dict,
) -> str:
    os.makedirs(params["dossier_output"], exist_ok=True)
    f = params["dossier_output"] + "ordonnancement_of.xlsx"

    for df in [df_cmd, df_det, df_crit, df_plan]:
        if df.empty:
            continue
        for col in df.columns:
            if df[col].dtype == "datetime64[ns]" or "date" in col.lower() or "Date" in col:
                try:
                    df[col] = pd.to_datetime(df[col], errors="coerce")
                    df[col] = df[col].dt.strftime("%d/%m/%Y").fillna("")
                except Exception:
                    pass

    with pd.ExcelWriter(f, engine="openpyxl") as w:
        if not df_cmd.empty:
            df_cmd.to_excel(w, sheet_name="Commandes", index=False)
        if not df_det.empty:
            df_det.to_excel(w, sheet_name="Detail Manquants", index=False)
        if not df_plan.empty:
            df_plan.to_excel(w, sheet_name="Plan de Charge", index=False)
        if not df_crit.empty:
            df_crit.to_excel(w, sheet_name="Composants Critiques", index=False)

    wb = load_workbook(f)

    FILL_ROW_VERT = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
    FILL_ROW_ORANGE = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
    FILL_ROW_ROUGE = PatternFill(start_color="FCE4EC", end_color="FCE4EC", fill_type="solid")
    ROW_FILLS = {"VERT": FILL_ROW_VERT, "ORANGE": FILL_ROW_ORANGE, "ROUGE": FILL_ROW_ROUGE}

    for sn in ["Commandes", "Detail Manquants", "Plan de Charge", "Composants Critiques"]:
        if sn not in wb.sheetnames:
            continue
        ws = wb[sn]
        headers = {ws.cell(row=1, column=c).value: c for c in range(1, ws.max_column + 1)}
        feu_cols = [h for h in headers if isinstance(h, str) and "Feu" in h]

        if sn == "Commandes" and "Feu Ligne" in headers:
            col_fl = headers["Feu Ligne"]
            col_tf = headers.get("Type Flux")
            for row in range(2, ws.max_row + 1):
                feu_val = ws.cell(row=row, column=col_fl).value
                is_se = col_tf and ws.cell(row=row, column=col_tf).value == "sous-ensemble"
                if not is_se and feu_val in ROW_FILLS:
                    for col in range(1, ws.max_column + 1):
                        ws.cell(row=row, column=col).fill = ROW_FILLS[feu_val]

        for row in range(2, ws.max_row + 1):
            for fc in feu_cols:
                _style_feu(ws, row, headers[fc], ws.cell(row=row, column=headers[fc]).value)

        if sn == "Commandes" and "Type Flux" in headers and "Designation" in headers:
            ct = headers["Type Flux"]
            cd = headers["Designation"]
            ws.sheet_properties.outlinePr = Outline(summaryBelow=False)
            for row in range(2, ws.max_row + 1):
                v = ws.cell(row=row, column=ct).value
                if v and str(v) == "sous-ensemble":
                    des = str(ws.cell(row=row, column=cd).value or "")
                    depth = 0
                    for ch in des:
                        if ch == ">":
                            depth += 1
                        else:
                            break
                    depth = max(1, depth // 2)
                    ws.row_dimensions[row].outline_level = depth
                    ws.row_dimensions[row].hidden = False
                    for col in range(1, ws.max_column + 1):
                        cell = ws.cell(row=row, column=col)
                        col_name = ws.cell(row=1, column=col).value
                        if col_name not in feu_cols:
                            cell.fill = FILL_SE
                        cell.font = FONT_SE

        if "Impact Livraison" in headers:
            ci = headers["Impact Livraison"]
            for row in range(2, ws.max_row + 1):
                v = ws.cell(row=row, column=ci).value
                if v and "RETARD" in str(v):
                    ws.cell(row=row, column=ci).fill = FILL_ROUGE
                    ws.cell(row=row, column=ci).font = FONT_ROUGE
                elif v == "OK":
                    ws.cell(row=row, column=ci).fill = FILL_VERT
                    ws.cell(row=row, column=ci).font = FONT_VERT
                elif v == "AUCUNE RECEPTION":
                    ws.cell(row=row, column=ci).fill = FILL_ROUGE
                    ws.cell(row=row, column=ci).font = FONT_ROUGE

        if sn == "Commandes" and "OF Associe" in headers and "Detail Manquants" in wb.sheetnames:
            co = headers["OF Associe"]
            ws_det = wb["Detail Manquants"]
            hd = {ws_det.cell(row=1, column=c).value: c for c in range(1, ws_det.max_column + 1)}
            col_mfg = hd.get("MFGNUM")
            if col_mfg:
                mfg_rows: Dict[str, int] = {}
                for dr_r in range(2, ws_det.max_row + 1):
                    v = ws_det.cell(row=dr_r, column=col_mfg).value
                    if v and v not in mfg_rows:
                        mfg_rows[v] = dr_r
                for row in range(2, ws.max_row + 1):
                    ov = ws.cell(row=row, column=co).value
                    if ov and str(ov).strip() in mfg_rows:
                        c_cell = ws.cell(row=row, column=co)
                        c_cell.hyperlink = f"#'Detail Manquants'!A{mfg_rows[str(ov).strip()]}"
                        c_cell.font = FONT_LINK

        if sn == "Plan de Charge":
            for col in range(1, ws.max_column + 1):
                h = ws.cell(row=1, column=col).value
                if h and str(h).endswith("(%)"):
                    for row in range(2, ws.max_row + 1):
                        v = ws.cell(row=row, column=col).value
                        if v is not None:
                            try:
                                val = float(v)
                                if val > 100:
                                    _style_feu(ws, row, col, "ROUGE")
                                elif val >= 80:
                                    _style_feu(ws, row, col, "ORANGE")
                                elif val > 0:
                                    _style_feu(ws, row, col, "VERT")
                            except Exception:
                                pass

        _fmt_sheet(ws, ws.max_column)

    wb.save(f)
    print(f"\n  Export: {f}")
    return f
