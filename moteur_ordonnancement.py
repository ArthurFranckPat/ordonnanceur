"""
Moteur d'ordonnancement OF - POC v3 PROD
Multi-niveaux recursif + allocation parametrique
Aereco / Sage X3 v12 - Donnees base production
"""

import pandas as pd
import numpy as np
from datetime import datetime, timedelta
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import Outline
import warnings, os

warnings.filterwarnings("ignore")

PARAMS = {
    "horizon_jours": 14,
    "mode_injection": "conservateur",
    "deduire_allocations": True,
    "allocation_sequentielle": True,
    "perimetre_chargement": "tous",        # "commandes" ou "tous"
    "capacite_heures_semaine": 70,        # 70h/semaine par poste
    "date_reference": datetime.now().replace(hour=0, minute=0, second=0, microsecond=0),
    "jours_ouvres_avant_expedition": 2,
    "separateur_csv": ";",
    "encoding": "latin-1",
    "dossier_data": "data/",
    "dossier_output": "output/",
}

FERMETURES_AERECO = []
STATUTS_NON_AFFERMIS = [2, 3]       # 2=Planifié, 3=Suggestion CBN
STATUTS_FERMES_LANCES = [1]          # 1=Ferme
PREFIXES_NON_BLOQUANTS = ["SF", "PF"]

# === CALENDRIER ===

def est_ferie_france(date):
    y=date.year
    fixes=[datetime(y,1,1),datetime(y,5,1),datetime(y,5,8),datetime(y,7,14),
           datetime(y,8,15),datetime(y,11,1),datetime(y,11,11),datetime(y,12,25)]
    a=y%19;b=y//100;c=y%100;d=b//4;e=b%4;f=(b+8)//25;g=(b-f+1)//3
    h=(19*a+b-d-g+15)%30;i=c//4;k=c%4;l=(32+2*e+2*i-h-k)%7
    m=(a+11*h+22*l)//451;mo=(h+l-7*m+114)//31;da=((h+l-7*m+114)%31)+1
    paques=datetime(y,mo,da)
    mobiles=[paques+timedelta(days=1),paques+timedelta(days=39),paques+timedelta(days=50)]
    return date in fixes or date in mobiles

def est_jour_ouvre(date, fermetures):
    return date.weekday()<5 and not est_ferie_france(date) and date not in fermetures

def soustraire_jours_ouvres(date, nb, fermetures):
    if pd.isna(date): return pd.NaT
    if isinstance(date, pd.Timestamp): date=date.to_pydatetime()
    r=date; n=nb
    while n>0:
        r-=timedelta(days=1)
        if est_jour_ouvre(r, fermetures): n-=1
    return r

def to_python_datetime(value):
    """Convertit une valeur pandas/OpenPyXL en datetime Python, sinon None."""
    if value is None or pd.isna(value):
        return None
    if isinstance(value, pd.Timestamp):
        return value.to_pydatetime()
    if isinstance(value, datetime):
        return value
    return None

# === CHARGEMENT ===

def to_num(series):
    if series.dtype==object:
        series=series.astype(str).str.strip().str.replace(' ','').str.replace(',','.')
    return pd.to_numeric(series, errors='coerce').fillna(0)

def lire_csv_normalise(path, sep, enc, colonnes_attendues, skipinitialspace=False):
    """Charge un CSV, retire les colonnes techniques et verifie le schema attendu."""
    df=pd.read_csv(path, sep=sep, encoding=enc, skipinitialspace=skipinitialspace)
    df=df.loc[:,~df.columns.str.startswith("Unnamed")]
    if len(df.columns)!=len(colonnes_attendues):
        raise ValueError(
            f"Schema inattendu pour {path}: {len(df.columns)} colonnes detectees "
            f"({list(df.columns)}) au lieu de {len(colonnes_attendues)} {colonnes_attendues}"
        )
    df.columns=colonnes_attendues
    return df

def charger_donnees(params):
    sep=params["separateur_csv"]; enc=params["encoding"]; d=params["dossier_data"]
    dfs={}

    # --- articles.csv (5 cols + trailing sep) ---
    # ARTICLE;DESCRIPTION;CATEGORIE;TYPE_APPRO;DELAI_REAPPRO;
    art=lire_csv_normalise(
        d+"articles.csv", sep, enc,
        ["itmref","designation","categorie","type_appro","delai"],
        skipinitialspace=True
    )
    art["delai"]=to_num(art["delai"])
    art["itmref"]=art["itmref"].astype(str).str.strip()
    art["categorie"]=art["categorie"].astype(str).str.strip()
    art["type_appro"]=art["type_appro"].astype(str).str.strip()
    dfs["articles"]=art

    # --- stock.csv (4 cols + trailing sep) ---
    # ARTICLE;STOCK_PHYSIQUE;STOCK_ALLOUE;STOCK_BLOQUE;
    stk=lire_csv_normalise(
        d+"stock.csv", sep, enc,
        ["itmref","stock_physique","stock_alloue","stock_bloque"],
        skipinitialspace=True
    )
    stk["itmref"]=stk["itmref"].astype(str).str.strip()
    for c in ["stock_physique","stock_alloue","stock_bloque"]: stk[c]=to_num(stk[c])
    dfs["stock"]=stk

    # --- commandes_clients.csv (11 cols + trailing sep) ---
    # NUM_COMMANDE;LIGNE_COMMANDE;CODE_CLIENT;NOM_CLIENT;ARTICLE;DESCRIPTION;
    # QTE_COMMANDEE;QTE_RESTANTE;DATE_EXPEDITION_DEMANDEE;FLAG_CONTREMARQUE;OF_CONTREMARQUE;
    cmd=lire_csv_normalise(
        d+"commandes_clients.csv", sep, enc,
        ["sohnum","soplin","client_code","client_nom","itmref","designation",
         "qte_commandee","qte_restante","shidat","flag_contremarque","mfgnum_lie"]
    )
    cmd["itmref"]=cmd["itmref"].astype(str).str.strip()
    for c in ["qte_commandee","qte_restante"]: cmd[c]=to_num(cmd[c])
    cmd["shidat"]=pd.to_datetime(cmd["shidat"],format="%d/%m/%Y",errors="coerce")
    cmd["flag_contremarque"]=to_num(cmd["flag_contremarque"]).astype(int)
    cmd["mfgnum_lie"]=cmd["mfgnum_lie"].astype(str).str.strip()
    dfs["commandes"]=cmd

    # --- of_entetes.csv (9 cols + trailing sep) ---
    # NUM_OF;ARTICLE;DESCRIPTION;STATUT_NUM_OF;STATUT_TEXTE_OF;DATE_FIN;
    # QTE_A_FABRIQUER;QTE_FABRIQUEE;QTE_RESTANTE;
    of=lire_csv_normalise(
        d+"of_entetes.csv", sep, enc,
        ["mfgnum","itmref","designation","mfgsta","mfgsta_lib",
         "enddat","extqty","cplqty","qte_restante"]
    )
    of["itmref"]=of["itmref"].astype(str).str.strip()
    of["mfgnum"]=of["mfgnum"].astype(str).str.strip()
    for c in ["extqty","cplqty","qte_restante"]: of[c]=to_num(of[c])
    of["enddat"]=pd.to_datetime(of["enddat"],format="%d/%m/%Y",errors="coerce")
    dfs["of_entetes"]=of

    # --- of_composants.csv (5 cols + trailing sep) ---
    # NUM_OF;ARTICLE;DESCRIPTION;QUANTITE_RESTANTE;DATE_BESOIN_COMPOSANT;
    # Note: QUANTITE_RESTANTE = quantite requise par composant (alias SQL)
    comp=lire_csv_normalise(
        d+"of_composants.csv", sep, enc,
        ["mfgnum","composant","designation","qty_requise","dat_besoin"]
    )
    comp["mfgnum"]=comp["mfgnum"].astype(str).str.strip()
    comp["composant"]=comp["composant"].astype(str).str.strip()
    comp["qty_requise"]=to_num(comp["qty_requise"])
    comp["dat_besoin"]=pd.to_datetime(comp["dat_besoin"],format="%d/%m/%Y",errors="coerce")
    dfs["of_composants"]=comp

    # --- receptions_oa.csv (5 cols + trailing sep) ---
    # NUM_COMMANDE;ARTICLE;CODE_FOURNISSEUR;QUANTITE_RESTANTE;DATE_RECEPTION_PREVUE;
    oa=lire_csv_normalise(
        d+"receptions_oa.csv", sep, enc,
        ["pohnum","itmref","fournisseur","qte_restante","date_reception"]
    )
    oa["itmref"]=oa["itmref"].astype(str).str.strip()
    oa["qte_restante"]=to_num(oa["qte_restante"])
    oa["date_reception"]=pd.to_datetime(oa["date_reception"],format="%d/%m/%Y",errors="coerce")
    dfs["receptions_oa"]=oa

    # --- receptions_of : reconstruit depuis of_entetes (fermes avec qte_restante > 0) ---
    of_fermes = of[(of["mfgsta"].isin(STATUTS_FERMES_LANCES)) & (of["qte_restante"]>0)].copy()
    rof = of_fermes[["mfgnum","itmref","qte_restante","enddat"]].copy()
    dfs["receptions_of"]=rof

    # --- gammes.csv (4 cols + trailing sep) ---
    # ARTICLE;POSTE_CHARGE;LIBELLE_POSTE;CADENCE;
    gam=lire_csv_normalise(
        d+"gammes.csv", sep, enc,
        ["itmref","poste_charge","libelle_poste","cadence"],
        skipinitialspace=True
    )
    gam["itmref"]=gam["itmref"].astype(str).str.strip()
    gam["poste_charge"]=gam["poste_charge"].astype(str).str.strip()
    gam["cadence"]=gam["cadence"].astype(str).str.replace(",",".").str.strip()
    gam["cadence"]=pd.to_numeric(gam["cadence"],errors="coerce").fillna(0)
    dfs["gammes"]=gam

    return dfs

# === PREPARATION ===

def preparer_donnees(dfs, params):
    date_ref=params["date_reference"]; date_lim=date_ref+timedelta(days=params["horizon_jours"])
    cmd=dfs["commandes"].copy()
    cmd=cmd[cmd["shidat"].notna() & (cmd["shidat"]<=date_lim) & (cmd["qte_restante"]>0)].copy()
    cmd["date_besoin_prod"]=cmd["shidat"].apply(
        lambda d: soustraire_jours_ouvres(d, params["jours_ouvres_avant_expedition"], FERMETURES_AERECO))
    cmd["type_flux"]=cmd.apply(
        lambda r: "contremarque" if r["flag_contremarque"]==1 and str(r["mfgnum_lie"]).strip() not in ["","nan","None","NaN","0","0.0"]
        else "sur_stock", axis=1)
    cmd=cmd.sort_values("shidat").reset_index(drop=True)
    dfs["commandes"]=cmd

    stk=dfs["stock"].copy()
    if params["deduire_allocations"]:
        stk["dispo_instantane"]=(stk["stock_physique"]-stk["stock_alloue"]-stk["stock_bloque"]).clip(lower=0)
    else:
        stk["dispo_instantane"]=stk["stock_physique"]
    dfs["stock"]=stk

    comp=dfs["of_composants"].copy()
    articles=dfs["articles"][["itmref","type_appro","categorie"]].copy()
    comp=comp.merge(articles.rename(columns={"itmref":"composant"}), on="composant", how="left")
    comp["est_non_bloquant"]=comp["categorie"].apply(
        lambda c: any(str(c).upper().startswith(p) for p in PREFIXES_NON_BLOQUANTS)
        if pd.notna(c) and str(c).strip() not in ["","nan"] else False)
    comp["flag_bloquant"]=~comp["est_non_bloquant"]
    dfs["of_composants"]=comp

    of_ent=dfs["of_entetes"].copy()
    dfs["of_non_affermis"]=of_ent[of_ent["mfgsta"].isin(STATUTS_NON_AFFERMIS)].copy()
    dfs["of_fermes_lances"]=of_ent[of_ent["mfgsta"].isin(STATUTS_FERMES_LANCES)].copy()
    return dfs

# === ALLOCATION STOCK PF ===

def allouer_stock_commandes(dfs):
    cmd=dfs["commandes"].copy()
    stock_dispo=dfs["stock"].set_index("itmref")["dispo_instantane"].to_dict()
    cmd["couverture_stock"]=0.0; cmd["ecart_a_couvrir"]=0.0
    cmd["of_associe"]=""; cmd["source_couverture"]=""
    for idx in cmd.index:
        row=cmd.loc[idx]
        if row["type_flux"]=="contremarque":
            cmd.at[idx,"of_associe"]=str(row["mfgnum_lie"]).strip()
            cmd.at[idx,"source_couverture"]="contremarque"
            cmd.at[idx,"ecart_a_couvrir"]=row["qte_restante"]; continue
        art=row["itmref"]; besoin=row["qte_restante"]; dispo=stock_dispo.get(art,0)
        if dispo>=besoin:
            cmd.at[idx,"couverture_stock"]=besoin; cmd.at[idx,"ecart_a_couvrir"]=0
            cmd.at[idx,"source_couverture"]="stock"; stock_dispo[art]=dispo-besoin
        elif dispo>0:
            cmd.at[idx,"couverture_stock"]=dispo; cmd.at[idx,"ecart_a_couvrir"]=besoin-dispo
            cmd.at[idx,"source_couverture"]="stock_partiel"; stock_dispo[art]=0
        else:
            cmd.at[idx,"ecart_a_couvrir"]=besoin; cmd.at[idx,"source_couverture"]="aucun_stock"
    dfs["commandes"]=cmd; return cmd

# === APPARIEMENT ===

def apparier_commandes_of(dfs):
    cmd=dfs["commandes"].copy()
    # Chercher dans TOUS les OF actifs avec qte_restante > 0
    of_all=dfs["of_entetes"].copy()
    of_candidats=of_all[of_all["qte_restante"]>0].copy()
    of_attribues=set(cmd[cmd["type_flux"]=="contremarque"]["of_associe"].dropna().unique())
    for idx in cmd.index:
        row=cmd.loc[idx]
        if row["type_flux"]=="contremarque" or row["ecart_a_couvrir"]<=0: continue
        candidats=of_candidats[
            (of_candidats["itmref"]==row["itmref"]) &
            (~of_candidats["mfgnum"].isin(of_attribues))
        ].copy()
        if candidats.empty:
            lbl="stock_partiel+aucun_of" if cmd.at[idx,"source_couverture"]=="stock_partiel" else "aucun_approvisionnement"
            cmd.at[idx,"source_couverture"]=lbl; continue
        date_besoin=row["date_besoin_prod"]
        if pd.notna(date_besoin):
            candidats_a_date=candidats[candidats["enddat"].notna() & (candidats["enddat"]<=date_besoin)].copy()
            if candidats_a_date.empty:
                lbl="stock_partiel+aucun_of_a_date" if cmd.at[idx,"source_couverture"]=="stock_partiel" else "aucun_of_a_date"
                cmd.at[idx,"source_couverture"]=lbl; continue
            candidats_a_date["ecart_date"]=(date_besoin-candidats_a_date["enddat"]).dt.days
            meilleur=candidats_a_date.sort_values(["ecart_date","enddat"]).iloc[0]
        else: meilleur=candidats.iloc[0]
        cmd.at[idx,"of_associe"]=meilleur["mfgnum"]; of_attribues.add(meilleur["mfgnum"])
        lbl="stock_partiel+of" if cmd.at[idx,"source_couverture"]=="stock_partiel" else "of"
        cmd.at[idx,"source_couverture"]=lbl
    dfs["commandes"]=cmd; return cmd

# === FAISABILITE ===

def calc_dispo_proj(article, date_besoin, dfs, stock_w):
    dispo=stock_w.get(article,0)
    if pd.isna(date_besoin): return dispo
    oa=dfs["receptions_oa"]
    dispo+=oa[(oa["itmref"]==article)&(oa["date_reception"].notna())&(oa["date_reception"]<=date_besoin)]["qte_restante"].sum()
    rof=dfs["receptions_of"]
    dispo+=rof[(rof["itmref"]==article)&(rof["enddat"].notna())&(rof["enddat"]<=date_besoin)]["qte_restante"].sum()
    return dispo

def calc_date_couverture(article, qty_requise, dfs, stock_w):
    dispo=stock_w.get(article,0)
    if dispo>=qty_requise: return None
    receptions=[]
    oa=dfs["receptions_oa"]
    for _,r in oa[(oa["itmref"]==article)&(oa["qte_restante"]>0)&(oa["date_reception"].notna())].iterrows():
        receptions.append((r["date_reception"], r["qte_restante"]))
    rof=dfs["receptions_of"]
    for _,r in rof[(rof["itmref"]==article)&(rof["qte_restante"]>0)&(rof["enddat"].notna())].iterrows():
        receptions.append((r["enddat"], r["qte_restante"]))
    receptions.sort(key=lambda x:x[0])
    cumul=dispo
    for dt,qt in receptions:
        cumul+=qt
        if cumul>=qty_requise: return dt
    return None

def trouver_of_pour_article(article, dfs, of_deja_traites):
    of_all=dfs["of_entetes"]
    candidats=of_all[(of_all["itmref"]==article)&(~of_all["mfgnum"].isin(of_deja_traites))].copy()
    if candidats.empty: return None
    na=candidats[candidats["mfgsta"].isin(STATUTS_NON_AFFERMIS)]
    if not na.empty: return na.iloc[0]["mfgnum"]
    fl=candidats[candidats["mfgsta"].isin(STATUTS_FERMES_LANCES)]
    if not fl.empty: return fl.iloc[0]["mfgnum"]
    return candidats.iloc[0]["mfgnum"]

def eval_of(mfgnum, dfs, stock_w, of_deja_traites=None, profondeur=0, max_prof=10):
    if of_deja_traites is None: of_deja_traites=set()
    empty_result={"mfgnum":mfgnum,"nb_composants":0,"ratio_instantane":0.0,
                  "ratio_projete":0.0,"feu":"ROUGE","detail_composants":[],
                  "nb_manquants_instant":0,"date_premiere_faisabilite":None,
                  "lancable":False,"of_enfants":[]}
    if profondeur > max_prof: return empty_result

    comp=dfs["of_composants"]; c_of=comp[comp["mfgnum"]==mfgnum].copy()
    of_info=dfs["of_entetes"][dfs["of_entetes"]["mfgnum"]==mfgnum]
    if c_of.empty or of_info.empty: return empty_result

    of_row=of_info.iloc[0]; date_besoin=of_row["enddat"]
    details=[]; ni=0; np_=0; nt=len(c_of); nm_instant=0
    ni_bloq=0; nt_bloq=0; np_bloq=0
    date_fais_max=None; of_enfants_resultats=[]; of_deja_traites.add(mfgnum)

    for _,c in c_of.iterrows():
        art=c["composant"]; qr=c["qty_requise"]; non_bloq=c["est_non_bloquant"]
        di=stock_w.get(art,0); dp=calc_dispo_proj(art,date_besoin,dfs,stock_w)
        ci=di>=qr; cp=dp>=qr

        if ci: feu_comp="VERT"
        elif cp: feu_comp="ORANGE"; nm_instant+=1
        else: feu_comp="ROUGE"; nm_instant+=1
        if ci: ni+=1
        if cp: np_+=1
        if not non_bloq:
            nt_bloq+=1
            if ci: ni_bloq+=1
            if cp: np_bloq+=1

        dc=calc_date_couverture(art,qr,dfs,stock_w) if not ci else None
        if dc is not None:
            if date_fais_max is None or dc>date_fais_max: date_fais_max=dc
        ecart_instant=max(0,qr-di)

        # Prochaine reception
        prd=None; prs=""
        oa=dfs["receptions_oa"]
        oa_a=oa[(oa["itmref"]==art)&(oa["qte_restante"]>0)&(oa["date_reception"].notna())]
        if not oa_a.empty:
            bo=oa_a.sort_values("date_reception").iloc[0]
            prd=bo["date_reception"]; prs=f"OA {bo['pohnum']}"
        rof_df=dfs["receptions_of"]
        rf_a=rof_df[(rof_df["itmref"]==art)&(rof_df["qte_restante"]>0)&(rof_df["enddat"].notna())]
        if not rf_a.empty:
            br=rf_a.sort_values("enddat").iloc[0]
            if prd is None or br["enddat"]<prd: prd=br["enddat"]; prs=f"OF {br['mfgnum']}"

        # Recursif : SF/PF en rupture -> OF enfant
        of_enfant_num=None; of_enfant_result=None; date_dispo_sf=None; retard_sf=None
        if non_bloq and not ci:
            of_enfant_num=trouver_of_pour_article(art, dfs, of_deja_traites)
            if of_enfant_num:
                of_enfant_result=eval_of(of_enfant_num, dfs, stock_w, of_deja_traites, profondeur+1)
                of_enfants_resultats.append(of_enfant_result)
                of_enfant_info=dfs["of_entetes"][dfs["of_entetes"]["mfgnum"]==of_enfant_num]
                enddat_enfant=of_enfant_info.iloc[0]["enddat"] if not of_enfant_info.empty else None
                df_of_enfant=of_enfant_result["date_premiere_faisabilite"]
                dates_c=[d for d in [enddat_enfant, df_of_enfant] if d is not None and not (isinstance(d,float) and np.isnan(d))]
                if dates_c:
                    date_dispo_sf=max(dates_c)
                    dat_besoin_comp=c["dat_besoin"]
                    if pd.notna(dat_besoin_comp) and pd.notna(date_dispo_sf):
                        d1=to_python_datetime(date_dispo_sf)
                        d2=to_python_datetime(dat_besoin_comp)
                        if d1 is not None and d2 is not None:
                            delta=(d1-d2).days
                            retard_sf=delta if delta>0 else 0

        details.append({
            "mfgnum":mfgnum,"composant":art,"designation":c["designation"],
            "categorie":c.get("categorie",""),"bloquant":"Oui" if c["flag_bloquant"] else "Non",
            "qty_requise":qr,"dispo_instantane":di,"dispo_projete":round(dp,2),
            "ecart_instant":ecart_instant,"feu_composant":feu_comp,"date_couverture":dc,
            "date_prochaine_reception":prd,"source_reception":prs,
            "of_enfant":of_enfant_num,"date_dispo_sf":date_dispo_sf,"retard_sf_jours":retard_sf,
        })

    ri=round(ni/nt*100,1) if nt else 0.0; rp=round(np_/nt*100,1) if nt else 0.0
    if rp==100: feu="VERT" if ri==100 else "ORANGE"
    elif rp>0: feu="ORANGE"
    else: feu="ROUGE"
    lancable=(np_bloq==nt_bloq) if nt_bloq>0 else True

    return {"mfgnum":mfgnum,"nb_composants":nt,"ratio_instantane":ri,"ratio_projete":rp,
            "feu":feu,"detail_composants":details,"nb_manquants_instant":nm_instant,
            "date_premiere_faisabilite":date_fais_max,"lancable":lancable,
            "of_enfants":of_enfants_resultats}

def alloc_seq_composants(of_list, dfs, params):
    stock_w=dfs["stock"].set_index("itmref")["dispo_instantane"].to_dict()
    alloc=params["allocation_sequentielle"]
    rb={}
    for m in of_list: rb[m]=eval_of(m,dfs,stock_w.copy())["ratio_projete"]
    of_ent=dfs["of_entetes"].set_index("mfgnum")
    of_tries=sorted(of_list, key=lambda m: (-rb.get(m,0),
        of_ent.loc[m,"enddat"] if m in of_ent.index else pd.Timestamp.max))
    if alloc: sa=stock_w.copy()
    res={}
    for m in of_tries:
        sw = sa.copy() if alloc else stock_w.copy()
        r=eval_of(m, dfs, sw, of_deja_traites=set())
        if alloc:
            comp=dfs["of_composants"]; cs=comp[comp["mfgnum"]==m]
            for _,c in cs.iterrows():
                if c["flag_bloquant"]:
                    a=c["composant"]; q=c["qty_requise"]; d=sa.get(a,0)
                    if d>=q: sa[a]=d-q
            for enf in r.get("of_enfants",[]):
                comp_enf=comp[comp["mfgnum"]==enf["mfgnum"]]
                for _,ce in comp_enf.iterrows():
                    if ce["flag_bloquant"]:
                        a=ce["composant"]; q=ce["qty_requise"]; d=sa.get(a,0)
                        if d>=q: sa[a]=d-q
        res[m]=r
    return res

# === CHARGEMENT CAPACITAIRE ===

def generer_jours_ouvres(date_debut, date_fin, fermetures):
    """Genere la liste des jours ouvres entre deux dates."""
    jours=[]
    d=date_debut
    while d<=date_fin:
        if est_jour_ouvre(d, fermetures): jours.append(d)
        d+=timedelta(days=1)
    return jours

def dernier_jour_ouvre_avant_ou_egal(date_cible, fermetures):
    """Retourne le dernier jour ouvre inferieur ou egal a la date cible."""
    if isinstance(date_cible, pd.Timestamp):
        date_cible = date_cible.to_pydatetime()
    d = date_cible
    while not est_jour_ouvre(d, fermetures):
        d -= timedelta(days=1)
    return d

def jour_ouvre_precedent(date_cible, fermetures):
    """Retourne le jour ouvre strictement precedent."""
    if isinstance(date_cible, pd.Timestamp):
        date_cible = date_cible.to_pydatetime()
    d = date_cible - timedelta(days=1)
    while not est_jour_ouvre(d, fermetures):
        d -= timedelta(days=1)
    return d

def calculer_plan_charge(dfs, resultats_of, params):
    """
    Construit le plan de charge par poste x jour.
    Backscheduling : charge placee au plus tard depuis date fin OF.
    """
    gammes=dfs["gammes"]
    of_ent=dfs["of_entetes"]
    date_ref=params["date_reference"]
    horizon=params["horizon_jours"]
    capa_jour=params["capacite_heures_semaine"]/5.0  # 70h/sem = 14h/jour
    perimetre=params["perimetre_chargement"]

    # Determiner les OF a charger
    if perimetre=="commandes":
        # OF lies aux commandes + leurs enfants
        of_a_charger=set()
        for on,r in resultats_of.items():
            of_a_charger.add(on)
            for enf in r.get("of_enfants",[]):
                of_a_charger.add(enf["mfgnum"])
        # Aussi les OF fermes evalues dans assembler (on les ajoutera apres)
    else:
        # Tous les OF actifs avec qte_restante > 0
        of_a_charger=set(of_ent[of_ent["qte_restante"]>0]["mfgnum"].values)

    # Gammes indexees par article
    gamme_idx=gammes.drop_duplicates(subset="itmref").set_index("itmref")

    # Jours ouvres de l'horizon affiche dans l'interface.
    # Le calcul de charge peut remonter avant cette date si un OF est deja en retard.
    date_lim=date_ref+timedelta(days=horizon)
    jours_ouvres_horizon=generer_jours_ouvres(date_ref, date_lim, FERMETURES_AERECO)
    if not jours_ouvres_horizon:
        return pd.DataFrame(), {}, pd.DataFrame()

    charge_map={}
    # Resultat par OF : {mfgnum: {"poste","charge_h","date_debut","date_fin","taux_moyen"}}
    of_capa={}

    for mfgnum in of_a_charger:
        of_info=of_ent[of_ent["mfgnum"]==mfgnum]
        if of_info.empty: continue
        oi=of_info.iloc[0]
        article=oi["itmref"]; qte=oi["qte_restante"]
        if qte<=0: continue

        # Chercher la gamme
        if article not in gamme_idx.index:
            of_capa[mfgnum]={"poste":"","libelle_poste":"","charge_h":0,
                             "date_debut":None,"date_fin":None,"alerte":"Gamme manquante"}
            continue
        gam=gamme_idx.loc[article]
        poste=gam["poste_charge"]; libelle=gam["libelle_poste"]; cadence=gam["cadence"]
        if cadence<=0:
            of_capa[mfgnum]={"poste":poste,"libelle_poste":libelle,"charge_h":0,
                             "date_debut":None,"date_fin":None,"alerte":"Cadence nulle"}
            continue

        charge_h=qte/cadence
        date_fin_of=oi["enddat"]
        if pd.isna(date_fin_of): date_fin_of=date_ref

        # Init poste dans charge_map
        if poste not in charge_map:
            charge_map[poste]={}

        # Backscheduling : placer la charge au plus tard,
        # y compris avant la date de reference si necessaire.
        jour=dernier_jour_ouvre_avant_ou_egal(date_fin_of, FERMETURES_AERECO)

        charge_restante=charge_h
        date_debut_charge=None
        date_fin_charge=None
        while charge_restante>0:
            if jour not in charge_map[poste]:
                charge_map[poste][jour]={"charge":0.0,"of_list":[],"allocations":[]}
            capa_restante=max(0, capa_jour - charge_map[poste][jour]["charge"])
            if capa_restante>0:
                charge_affectee=min(charge_restante, capa_restante)
                charge_map[poste][jour]["charge"]+=charge_affectee
                charge_map[poste][jour]["of_list"].append(mfgnum)
                charge_map[poste][jour]["allocations"].append({"mfgnum":mfgnum,"charge_h":charge_affectee})
                charge_restante-=charge_affectee
                date_debut_charge=jour
                if date_fin_charge is None: date_fin_charge=jour
            if charge_restante>0:
                jour=jour_ouvre_precedent(jour, FERMETURES_AERECO)

        of_capa[mfgnum]={"poste":poste,"libelle_poste":libelle,"charge_h":round(charge_h,2),
                         "date_debut":date_debut_charge,"date_fin":date_fin_charge,"alerte":""}

    jours_affiches=set(jours_ouvres_horizon)
    for poste_map in charge_map.values():
        jours_affiches.update(poste_map.keys())
    jours_ouvres=sorted(jours_affiches)

    # Construire le DataFrame plan de charge — MAILLE SEMAINE
    capa_sem=params["capacite_heures_semaine"]
    postes_tous=sorted(charge_map.keys())

    # Agreger charge jour -> semaine
    # Semaine = (annee, num_semaine)
    semaine_charge={}  # {(annee,sem): {poste: {"charge":0, "of_set":set()}}}
    for jour in jours_ouvres:
        iso=jour.isocalendar()
        sem_key=(iso[0], iso[1])
        if sem_key not in semaine_charge:
            semaine_charge[sem_key]={}
        for poste in postes_tous:
            if poste not in semaine_charge[sem_key]:
                semaine_charge[sem_key][poste]={"charge":0.0,"of_set":set()}
            ch=charge_map.get(poste,{}).get(jour,{}).get("charge",0)
            ofs=charge_map.get(poste,{}).get(jour,{}).get("of_list",[])
            semaine_charge[sem_key][poste]["charge"]+=ch
            semaine_charge[sem_key][poste]["of_set"].update(ofs)

    rows=[]
    for sem_key in sorted(semaine_charge.keys()):
        annee,num_sem=sem_key
        row={"Semaine":f"S{num_sem:02d}-{annee}","Capa (h/sem)":capa_sem}
        for poste in postes_tous:
            info_sem=semaine_charge[sem_key].get(poste,{"charge":0,"of_set":set()})
            ch=info_sem["charge"]
            taux=round(ch/capa_sem*100,1) if capa_sem>0 else 0
            nb_of=len(info_sem["of_set"])
            row[f"{poste} (h)"]=round(ch,1)
            row[f"{poste} (%)"]=taux
            row[f"{poste} (OF)"]=nb_of
        rows.append(row)
    df_plan=pd.DataFrame(rows)

    # Calculer le taux moyen par poste SEMAINE pour chaque OF
    for mfgnum,info in of_capa.items():
        if info["date_debut"] and info["date_fin"] and info["poste"]:
            poste=info["poste"]
            d1=info["date_debut"]; d2=info["date_fin"]
            jours_of=generer_jours_ouvres(d1, d2, FERMETURES_AERECO) if d1<=d2 else []
            semaines_of=set()
            for j in jours_of:
                iso=j.isocalendar()
                semaines_of.add((iso[0],iso[1]))
            if semaines_of:
                taux_list=[]
                for sk in semaines_of:
                    ch_sem=semaine_charge.get(sk,{}).get(poste,{}).get("charge",0)
                    taux_list.append(ch_sem/capa_sem*100 if capa_sem>0 else 0)
                info["taux_moyen"]=round(sum(taux_list)/len(taux_list),1)
            else:
                info["taux_moyen"]=0
        else:
            info["taux_moyen"]=0

    charge_segment_rows=[]
    heure_debut_journee=8.0
    for poste in postes_tous:
        for jour in jours_ouvres:
            allocs=charge_map.get(poste,{}).get(jour,{}).get("allocations",[])
            if not allocs:
                continue
            charge_totale_jour=sum(float(a["charge_h"]) for a in allocs)
            fenetre_visible_h=capa_jour if capa_jour>0 else 14.0
            ratio_affichage=min(1.0, fenetre_visible_h/charge_totale_jour) if charge_totale_jour>0 else 1.0
            cumul_h=0.0
            for idx,alloc in enumerate(allocs):
                mfgnum=alloc["mfgnum"]
                charge_h=float(alloc["charge_h"])
                iso=jour.isocalendar()
                sem_label=f"S{iso[1]:02d}-{iso[0]}"
                charge_h_visible=charge_h*ratio_affichage
                segment_start=jour + timedelta(hours=heure_debut_journee+cumul_h)
                segment_end=jour + timedelta(hours=heure_debut_journee+cumul_h+charge_h_visible)
                cumul_h+=charge_h_visible

                of_info=of_ent[of_ent["mfgnum"]==mfgnum]
                if of_info.empty:
                    article=""; designation=""; statut_num=None; statut_of=""; date_fin_of=pd.NaT; date_debut_of=pd.NaT
                else:
                    oi=of_info.iloc[0]
                    article=oi["itmref"]
                    designation=oi["designation"]
                    statut_num=oi["mfgsta"]
                    statut_of=oi["mfgsta_lib"]
                    date_fin_of=oi["enddat"]
                    date_debut_of=of_capa.get(mfgnum,{}).get("date_debut", pd.NaT)

                charge_segment_rows.append({
                    "Semaine":sem_label,
                    "Date":jour,
                    "Poste":poste,
                    "MFGNUM":mfgnum,
                    "Article":article,
                    "Designation":designation,
                    "Statut OF":statut_of,
                    "Statut Num":statut_num,
                    "Date debut OF":date_debut_of,
                    "Date fin OF":date_fin_of,
                    "Charge (h)":round(charge_h,2),
                    "Charge visible (h)":round(charge_h_visible,2),
                    "Segment debut":segment_start,
                    "Segment fin":segment_end,
                    "Ordre jour":idx,
                })

    if charge_segment_rows:
        df_charge_segments=pd.DataFrame(charge_segment_rows)
        df_charge_segments=df_charge_segments.sort_values(["Semaine","Poste","Date","Ordre jour","MFGNUM"]).reset_index(drop=True)
        df_charge_segments=df_charge_segments.drop(columns=["Ordre jour"])
    else:
        df_charge_segments=pd.DataFrame(columns=["Semaine","Date","Poste","MFGNUM","Article","Designation","Statut OF","Statut Num","Date debut OF","Date fin OF","Charge (h)","Charge visible (h)","Segment debut","Segment fin"])

    return df_plan, of_capa, df_charge_segments

# === ASSEMBLAGE ===

def collecter_enfants_inline(result, dfs, profondeur=0):
    rows=[]
    for enf in result.get("of_enfants",[]):
        of_info=dfs["of_entetes"][dfs["of_entetes"]["mfgnum"]==enf["mfgnum"]]
        if not of_info.empty:
            oi=of_info.iloc[0]
            article=oi["itmref"]; designation=oi["designation"]
            date_fin=oi["enddat"]; qte_fab=oi["extqty"]; qte_prod=oi["cplqty"]
        else:
            article=""; designation=""; date_fin=pd.NaT; qte_fab=0; qte_prod=0
        prefix=">>" * (profondeur+1) + " "
        rows.append({
            "article":article,"designation":prefix+designation,
            "qte_fab":qte_fab,"qte_prod":qte_prod,"date_fin":date_fin,
            "date_faisabilite":enf["date_premiere_faisabilite"],
            "type_flux":"sous-ensemble",
            "ratio_instant":enf["ratio_instantane"],"ratio_projete":enf["ratio_projete"],
            "feu":enf["feu"],"nb_manquants":enf["nb_manquants_instant"],
            "lancable":"OUI" if enf["lancable"] else "NON","mfgnum":enf["mfgnum"],
        })
        rows.extend(collecter_enfants_inline(enf, dfs, profondeur+1))
    return rows

def assembler(dfs, resultats_of, of_capa=None):
    if of_capa is None: of_capa={}
    cmd=dfs["commandes"].copy()
    cmd["ratio_instant_of"]=np.nan; cmd["ratio_projete_of"]=np.nan
    cmd["feu_of"]=None; cmd["nb_manquants_of"]=np.nan
    cmd["feu_matiere"]=None; cmd["alerte"]=None
    cmd["date_faisabilite"]=pd.NaT; cmd["lancable"]=None; cmd["nb_of_enfants"]=0
    cmd["poste_charge"]=None; cmd["feu_capacite"]=None; cmd["feu_ligne"]=None

    def feu_capa_of(mfgnum):
        info=of_capa.get(mfgnum,{})
        if not info or info.get("alerte"): return None
        t=info.get("taux_moyen",0)
        if t>100: return "ROUGE"
        elif t>=80: return "ORANGE"
        else: return "VERT"

    def pire_feu(f1,f2):
        ordre={"ROUGE":0,"ORANGE":1,"VERT":2}
        if f1 is None: return f2
        if f2 is None: return f1
        return f1 if ordre.get(f1,2)<=ordre.get(f2,2) else f2

    of_dates=dfs["of_entetes"][["mfgnum","enddat"]].drop_duplicates(subset="mfgnum")
    of_dates_idx=of_dates.set_index("mfgnum") if not of_dates.empty else pd.DataFrame()

    for idx in cmd.index:
        row=cmd.loc[idx]; on=str(row["of_associe"]).strip()
        date_besoin_prod=row.get("date_besoin_prod", pd.NaT)
        if row["source_couverture"]=="stock":
            cmd.at[idx,"feu_matiere"]="VERT"; cmd.at[idx,"feu_ligne"]="VERT"
            cmd.at[idx,"lancable"]="-"; continue
        if "aucun" in str(row["source_couverture"]):
            cmd.at[idx,"feu_matiere"]="ROUGE"; cmd.at[idx,"feu_ligne"]="ROUGE"
            cmd.at[idx,"alerte"]="Aucun OF trouve"; cmd.at[idx,"lancable"]="-"; continue
        if on in resultats_of:
            r=resultats_of[on]
            cmd.at[idx,"ratio_instant_of"]=r["ratio_instantane"]
            cmd.at[idx,"ratio_projete_of"]=r["ratio_projete"]
            cmd.at[idx,"feu_of"]=r["feu"]
            cmd.at[idx,"nb_manquants_of"]=r["nb_manquants_instant"]
            cmd.at[idx,"feu_matiere"]=r["feu"]
            cmd.at[idx,"lancable"]="OUI" if r["lancable"] else "NON"
            if r["date_premiere_faisabilite"] is not None:
                cmd.at[idx,"date_faisabilite"]=r["date_premiere_faisabilite"]
            cmd.at[idx,"nb_of_enfants"]=len(r.get("of_enfants",[]))
        elif on and on in dfs["of_fermes_lances"]["mfgnum"].values:
            r_ferme=eval_of(on, dfs, dfs["stock"].set_index("itmref")["dispo_instantane"].to_dict())
            cmd.at[idx,"ratio_instant_of"]=r_ferme["ratio_instantane"]
            cmd.at[idx,"ratio_projete_of"]=r_ferme["ratio_projete"]
            cmd.at[idx,"feu_of"]=r_ferme["feu"]
            cmd.at[idx,"nb_manquants_of"]=r_ferme["nb_manquants_instant"]
            cmd.at[idx,"feu_matiere"]=r_ferme["feu"]
            cmd.at[idx,"lancable"]="FERME"
            if r_ferme["date_premiere_faisabilite"] is not None:
                cmd.at[idx,"date_faisabilite"]=r_ferme["date_premiere_faisabilite"]
            if r_ferme["nb_manquants_instant"]>0:
                resultats_of[on]=r_ferme
        else:
            cmd.at[idx,"feu_matiere"]="ROUGE"; cmd.at[idx,"feu_ligne"]="ROUGE"
            cmd.at[idx,"alerte"]=f"OF {on} introuvable"; cmd.at[idx,"lancable"]="-"; continue

        if on and not of_dates_idx.empty and on in of_dates_idx.index and pd.notna(date_besoin_prod):
            date_fin_of=of_dates_idx.loc[on,"enddat"]
            date_fin_of_dt=to_python_datetime(date_fin_of)
            date_besoin_prod_dt=to_python_datetime(date_besoin_prod)
            if date_fin_of_dt is not None and date_besoin_prod_dt is not None and date_fin_of_dt>date_besoin_prod_dt:
                retard_j=(date_fin_of_dt-date_besoin_prod_dt).days
                alerte_retard=f"OF apres besoin commande (+{retard_j}j)"
                alerte_existante=cmd.at[idx,"alerte"]
                cmd.at[idx,"alerte"]=alerte_retard if pd.isna(alerte_existante) or not alerte_existante else f"{alerte_existante} | {alerte_retard}"
                if row["type_flux"]=="contremarque":
                    cmd.at[idx,"feu_ligne"]="ROUGE"

        # Capacite
        fc=feu_capa_of(on)
        capa_info=of_capa.get(on,{})
        cmd.at[idx,"poste_charge"]=capa_info.get("poste","")
        cmd.at[idx,"feu_capacite"]=fc
        if capa_info.get("alerte"):
            cmd.at[idx,"alerte"]=capa_info["alerte"]
        # Feu global = pire des deux
        cmd.at[idx,"feu_ligne"]=pire_feu(cmd.at[idx,"feu_matiere"], fc)

    # Construire onglet Commandes avec sous-ensembles intercales
    colonnes=["N Commande","Ligne","Code Client","Nom Client","Article","Designation",
        "Qte Commandee","Qte Restante","Date Expedition","Date Besoin Prod","Type Flux",
        "Couvert par Stock","Ecart a Couvrir","OF Associe","Source Couverture",
        "Ratio Instant OF (%)","Ratio Projete OF (%)","Feu Matiere","Feu Capacite","Feu Ligne",
        "Poste Charge","Nb Manquants (instant)","Lancable","Date 1ere Faisabilite",
        "Nb OF Enfants","Alerte"]
    all_rows=[]
    for idx in cmd.index:
        row=cmd.loc[idx]
        all_rows.append({
            "N Commande":row["sohnum"],"Ligne":row["soplin"],
            "Code Client":row["client_code"],"Nom Client":row["client_nom"],
            "Article":row["itmref"],"Designation":row["designation"],
            "Qte Commandee":row["qte_commandee"],"Qte Restante":row["qte_restante"],
            "Date Expedition":row["shidat"],"Date Besoin Prod":row["date_besoin_prod"],
            "Type Flux":row["type_flux"],
            "Couvert par Stock":row["couverture_stock"],"Ecart a Couvrir":row["ecart_a_couvrir"],
            "OF Associe":row["of_associe"],"Source Couverture":row["source_couverture"],
            "Ratio Instant OF (%)":row["ratio_instant_of"],
            "Ratio Projete OF (%)":row["ratio_projete_of"],
            "Feu Matiere":row["feu_matiere"],"Feu Capacite":row["feu_capacite"],
            "Feu Ligne":row["feu_ligne"],
            "Poste Charge":row["poste_charge"],
            "Nb Manquants (instant)":row["nb_manquants_of"],
            "Lancable":row["lancable"],"Date 1ere Faisabilite":row["date_faisabilite"],
            "Nb OF Enfants":row["nb_of_enfants"],"Alerte":row["alerte"],
        })
        on=str(row["of_associe"]).strip()
        if on in resultats_of:
            enfants=collecter_enfants_inline(resultats_of[on], dfs)
            for enf in enfants:
                enf_capa=of_capa.get(enf["mfgnum"],{})
                all_rows.append({
                    "N Commande":"","Ligne":"","Code Client":"","Nom Client":"",
                    "Article":enf["article"],"Designation":enf["designation"],
                    "Qte Commandee":enf["qte_fab"],"Qte Restante":enf["qte_prod"],
                    "Date Expedition":enf["date_fin"],"Date Besoin Prod":"",
                    "Type Flux":"sous-ensemble","Couvert par Stock":"","Ecart a Couvrir":"",
                    "OF Associe":enf["mfgnum"],"Source Couverture":"",
                    "Ratio Instant OF (%)":enf["ratio_instant"],
                    "Ratio Projete OF (%)":enf["ratio_projete"],
                    "Feu Matiere":enf["feu"],"Feu Capacite":feu_capa_of(enf["mfgnum"]),
                    "Feu Ligne":"",
                    "Poste Charge":enf_capa.get("poste",""),
                    "Nb Manquants (instant)":enf["nb_manquants"],
                    "Lancable":enf["lancable"],"Date 1ere Faisabilite":enf["date_faisabilite"],
                    "Nb OF Enfants":"","Alerte":enf_capa.get("alerte",""),
                })
    df_cmd=pd.DataFrame(all_rows, columns=colonnes)

    # --- Detail manquants ---
    dr=[]
    for on,r in resultats_of.items():
        cl=cmd[cmd["of_associe"]==on]
        soh=cl["sohnum"].iloc[0] if not cl.empty else ""
        date_besoin_prod=cl["date_besoin_prod"].iloc[0] if not cl.empty else pd.NaT
        for d in r["detail_composants"]:
            if d["dispo_instantane"]>=d["qty_requise"]: continue
            retard_flag=""; retard_j=None
            if pd.notna(d["date_prochaine_reception"]) and pd.notna(date_besoin_prod):
                dr_=to_python_datetime(d["date_prochaine_reception"])
                dbp=to_python_datetime(date_besoin_prod)
                if dr_ is not None and dbp is not None:
                    delta=(dr_-dbp).days
                    if delta>0: retard_j=delta; retard_flag=f"RETARD {delta}j"
                    else: retard_flag="OK"
            elif d["date_prochaine_reception"] is None or pd.isna(d.get("date_prochaine_reception")):
                retard_flag="AUCUNE RECEPTION"
            of_enf_info=""
            if d.get("of_enfant"):
                of_enf_info=d["of_enfant"]
                if d.get("retard_sf_jours") and d["retard_sf_jours"]>0:
                    of_enf_info+=f" (RETARD {d['retard_sf_jours']}j)"
                elif d.get("date_dispo_sf"): of_enf_info+=" (OK)"
            dr.append({"N Commande":soh,"MFGNUM":on,
                "Composant":d["composant"],"Designation":d["designation"],
                "Categorie":d["categorie"],"Bloquant":d["bloquant"],
                "Qte Requise":d["qty_requise"],"Dispo Instantane":d["dispo_instantane"],
                "Dispo Projete":d["dispo_projete"],"Ecart (instant)":d["ecart_instant"],
                "Feu":d["feu_composant"],
                "Date Proch Reception":d["date_prochaine_reception"],
                "Source Reception":d["source_reception"],
                "Date Couverture":d["date_couverture"],
                "Impact Livraison":retard_flag,"Retard (jours)":retard_j,
                "OF Enfant (SF/PF)":of_enf_info,"Date Dispo SF":d.get("date_dispo_sf"),
            })
    df_det=pd.DataFrame(dr) if dr else pd.DataFrame()

    # --- Composants critiques ---
    ad=[]
    for on,r in resultats_of.items():
        for d in r["detail_composants"]:
            if d["bloquant"]=="Oui": ad.append(d)
        for enf in r.get("of_enfants",[]):
            for d in enf["detail_composants"]:
                if d["bloquant"]=="Oui": ad.append(d)
    if ad:
        da=pd.DataFrame(ad)
        cr=da.groupby("composant").agg(
            designation=("designation","first"),categorie=("categorie","first"),
            dispo_instantane=("dispo_instantane","first"),
            besoin_total=("qty_requise","sum"),nb_of=("mfgnum","nunique"),
            of_list=("mfgnum",lambda x:", ".join(sorted(x.unique())))).reset_index()
        cr["taux_couv"]=(cr["dispo_instantane"]/cr["besoin_total"].replace(0,np.nan)*100).round(1)
        def cnt_cmd(ol):
            return len(set(s for o in ol.split(", ") for s in cmd[cmd["of_associe"]==o]["sohnum"].unique()))
        cr["nb_cmd"]=cr["of_list"].apply(cnt_cmd)
        cr=cr.sort_values("taux_couv").reset_index(drop=True)
        cr.columns=["Composant","Designation","Categorie","Stock Dispo","Besoin Total",
                     "Nb OF","OF Concernes","Taux Couverture (%)","Nb Commandes"]
    else: cr=pd.DataFrame()
    return df_cmd, df_det, cr

# === ORCHESTRATION ===

def run(params=None):
    if params is None: params=PARAMS
    alloc_str="OUI" if params["allocation_sequentielle"] else "NON"
    print("="*70)
    print(f"  ORDONNANCEMENT OF v3 PROD — Aereco")
    print(f"  Date ref: {params['date_reference'].strftime('%d/%m/%Y')} | Horizon {params['horizon_jours']}j | Alloc seq: {alloc_str}")
    print("="*70)

    print("\n[1/9] Chargement...")
    dfs=charger_donnees(params)
    print(f"  {len(dfs['commandes'])} cmd | {len(dfs['of_entetes'])} OF | {len(dfs['of_composants'])} comp")
    print(f"  {len(dfs['stock'])} stk | {len(dfs['receptions_oa'])} OA | {len(dfs['receptions_of'])} rec OF | {len(dfs['gammes'])} gammes")

    print("\n[2/9] Preparation...")
    dfs=preparer_donnees(dfs,params)
    cmd=dfs["commandes"]
    print(f"  {len(cmd)} lignes horizon | {len(cmd[cmd['type_flux']=='contremarque'])} ctm | {len(cmd[cmd['type_flux']=='sur_stock'])} stock")
    print(f"  {len(dfs['of_non_affermis'])} OF non affermis | {len(dfs['of_fermes_lances'])} OF fermes")

    print("\n[3/9] Allocation stock PF...")
    cmd=allouer_stock_commandes(dfs)
    print(f"  {len(cmd[cmd['source_couverture']=='stock'])} couvertes par stock")

    print("\n[4/9] Appariement cmd->OF...")
    cmd=apparier_commandes_of(dfs)
    print(f"  {len(cmd[cmd['source_couverture'].str.contains('aucun',na=False)])} sans OF")

    print("\n[5/9] Identification OF...")
    of_list=list(cmd[cmd["of_associe"].astype(str).str.strip().ne("")]["of_associe"].unique())
    of_na_set=set(dfs["of_non_affermis"]["mfgnum"].values)
    of_eval=[o for o in of_list if o in of_na_set]
    of_ferme=[o for o in of_list if o not in of_na_set]
    print(f"  {len(of_eval)} OF non affermis | {len(of_ferme)} OF fermes (seront aussi evalues)")

    print("\n[6/9] Faisabilite + multi-niveaux recursif...")
    if of_eval:
        res_of=alloc_seq_composants(of_eval,dfs,params)
        nb_enf=sum(len(r.get("of_enfants",[])) for r in res_of.values())
        print(f"  {len(res_of)} OF parents | {nb_enf} OF enfants decouverts")
    else:
        res_of={}; print("  Aucun OF non affermi")

    print("\n[7/9] Chargement capacitaire (backscheduling)...")
    df_plan,of_capa,df_charge_segments=calculer_plan_charge(dfs,res_of,params)
    nb_gammes=len([m for m,i in of_capa.items() if not i.get("alerte")])
    nb_sans_gamme=len([m for m,i in of_capa.items() if i.get("alerte")=="Gamme manquante"])
    print(f"  {nb_gammes} OF avec gamme | {nb_sans_gamme} sans gamme")
    if not df_plan.empty:
        pct_cols=[c for c in df_plan.columns if c.endswith("(%)")]
        surcharges=0
        for c in pct_cols:
            surcharges+=(df_plan[c]>100).sum()
        print(f"  {len(pct_cols)} postes charges | {surcharges} semaines en surcharge (>100%)")

    print("\n[8/9] Assemblage (+ evaluation OF fermes)...")
    df_cmd,df_det,df_crit=assembler(dfs,res_of,of_capa)

    if not df_charge_segments.empty:
        cmd_for_order=dfs["commandes"].copy()
        cmd_for_order=cmd_for_order[cmd_for_order["of_associe"].astype(str).str.strip().ne("")].copy()
        if not cmd_for_order.empty:
            of_order_info=cmd_for_order.groupby("of_associe", as_index=False).agg({
                "date_besoin_prod":"min",
                "shidat":"min",
            }).rename(columns={
                "of_associe":"MFGNUM",
                "date_besoin_prod":"Date besoin cmd",
                "shidat":"Date expedition cmd",
            })
            df_charge_segments=df_charge_segments.merge(of_order_info, on="MFGNUM", how="left")
        else:
            df_charge_segments["Date besoin cmd"]=pd.NaT
            df_charge_segments["Date expedition cmd"]=pd.NaT

    print("\n[9/9] Resultats")
    print("="*70)
    if not df_cmd.empty:
        nv=len(df_cmd[df_cmd["Feu Ligne"]=="VERT"])
        no=len(df_cmd[df_cmd["Feu Ligne"]=="ORANGE"])
        nr=len(df_cmd[df_cmd["Feu Ligne"]=="ROUGE"])
        nse=len(df_cmd[df_cmd["Type Flux"]=="sous-ensemble"])
        print(f"  VERT {nv} | ORANGE {no} | ROUGE {nr} | Sous-ensembles {nse}")
    if not df_crit.empty:
        print(f"\n  Top composants critiques:")
        for _,r in df_crit.head(10).iterrows():
            print(f"    {r['Composant']:15s} couv {r['Taux Couverture (%)']:6.1f}% — {r['Nb OF']} OF, {r['Nb Commandes']} cmd")
    return df_cmd,df_det,df_plan,df_crit,df_charge_segments

# === EXPORT EXCEL ===

FILL_VERT=PatternFill(start_color="C6EFCE",end_color="C6EFCE",fill_type="solid")
FILL_ORANGE=PatternFill(start_color="FFEB9C",end_color="FFEB9C",fill_type="solid")
FILL_ROUGE=PatternFill(start_color="FFC7CE",end_color="FFC7CE",fill_type="solid")
FILL_SE=PatternFill(start_color="D9E2F3",end_color="D9E2F3",fill_type="solid")
FONT_VERT=Font(color="006100",bold=True); FONT_ORANGE=Font(color="9C5700",bold=True)
FONT_ROUGE=Font(color="9C0006",bold=True); FONT_SE=Font(color="2F5496",italic=True)
FONT_HEADER=Font(bold=True,color="FFFFFF",size=10)
FILL_HEADER=PatternFill(start_color="4472C4",end_color="4472C4",fill_type="solid")
FONT_LINK=Font(color="0563C1",underline="single")
THIN_BORDER=Border(left=Side(style='thin'),right=Side(style='thin'),
                   top=Side(style='thin'),bottom=Side(style='thin'))
FEU_STYLES={"VERT":(FILL_VERT,FONT_VERT),"ORANGE":(FILL_ORANGE,FONT_ORANGE),"ROUGE":(FILL_ROUGE,FONT_ROUGE)}

def style_feu(ws,row,col,val):
    if val in FEU_STYLES:
        f,fo=FEU_STYLES[val]; ws.cell(row=row,column=col).fill=f; ws.cell(row=row,column=col).font=fo

def fmt_sheet(ws,nc):
    for col in range(1,nc+1):
        c=ws.cell(row=1,column=col); c.font=FONT_HEADER; c.fill=FILL_HEADER
        c.alignment=Alignment(horizontal="center",wrap_text=True)
    for col in range(1,nc+1):
        ml=0
        for row in range(1,min(ws.max_row+1,100)):
            v=ws.cell(row=row,column=col).value
            if v: ml=max(ml,len(str(v)))
        ws.column_dimensions[get_column_letter(col)].width=min(ml+3,30)
    for row in range(1,ws.max_row+1):
        for col in range(1,nc+1):
            ws.cell(row=row,column=col).border=THIN_BORDER

def export_xl(df_cmd,df_det,df_plan,df_crit,params=None):
    if params is None: params=PARAMS
    os.makedirs(params["dossier_output"],exist_ok=True)
    f=params["dossier_output"]+"ordonnancement_of.xlsx"

    # Formater les dates en DD/MM/YYYY avant export
    for df in [df_cmd, df_det, df_crit, df_plan]:
        if df.empty: continue
        for col in df.columns:
            if df[col].dtype == 'datetime64[ns]' or 'date' in col.lower() or 'Date' in col:
                try:
                    df[col] = pd.to_datetime(df[col], errors='coerce')
                    df[col] = df[col].dt.strftime("%d/%m/%Y").fillna("")
                except: pass

    with pd.ExcelWriter(f,engine="openpyxl") as w:
        if not df_cmd.empty: df_cmd.to_excel(w,sheet_name="Commandes",index=False)
        if not df_det.empty: df_det.to_excel(w,sheet_name="Detail Manquants",index=False)
        if not df_plan.empty: df_plan.to_excel(w,sheet_name="Plan de Charge",index=False)
        if not df_crit.empty: df_crit.to_excel(w,sheet_name="Composants Critiques",index=False)

    from openpyxl import load_workbook
    wb=load_workbook(f)

    # Fills legers pour les lignes entieres
    FILL_ROW_VERT=PatternFill(start_color="E2EFDA",end_color="E2EFDA",fill_type="solid")
    FILL_ROW_ORANGE=PatternFill(start_color="FFF2CC",end_color="FFF2CC",fill_type="solid")
    FILL_ROW_ROUGE=PatternFill(start_color="FCE4EC",end_color="FCE4EC",fill_type="solid")
    ROW_FILLS={"VERT":FILL_ROW_VERT,"ORANGE":FILL_ROW_ORANGE,"ROUGE":FILL_ROW_ROUGE}

    for sn in ["Commandes","Detail Manquants","Plan de Charge","Composants Critiques"]:
        if sn not in wb.sheetnames: continue
        ws=wb[sn]; headers={ws.cell(row=1,column=c).value:c for c in range(1,ws.max_column+1)}
        feu_cols=[h for h in headers if isinstance(h, str) and "Feu" in h]

        # --- Coloration lignes entieres selon Feu Ligne ---
        if sn=="Commandes" and "Feu Ligne" in headers:
            col_fl=headers["Feu Ligne"]
            col_tf=headers.get("Type Flux")
            for row in range(2,ws.max_row+1):
                feu_val=ws.cell(row=row,column=col_fl).value
                is_se = col_tf and ws.cell(row=row,column=col_tf).value=="sous-ensemble"
                if not is_se and feu_val in ROW_FILLS:
                    for col in range(1,ws.max_column+1):
                        ws.cell(row=row,column=col).fill=ROW_FILLS[feu_val]

        # --- Coloration cellules Feu (plus intense, par dessus) ---
        for row in range(2,ws.max_row+1):
            for fc in feu_cols:
                style_feu(ws,row,headers[fc],ws.cell(row=row,column=headers[fc]).value)

        # --- Sous-ensemble: accordeon + style (ecrase la couleur ligne) ---
        if sn=="Commandes" and "Type Flux" in headers and "Designation" in headers:
            ct=headers["Type Flux"]; cd=headers["Designation"]
            ws.sheet_properties.outlinePr=Outline(summaryBelow=False)
            for row in range(2,ws.max_row+1):
                v=ws.cell(row=row,column=ct).value
                if v and str(v)=="sous-ensemble":
                    des=str(ws.cell(row=row,column=cd).value or "")
                    depth=0
                    for ch in des:
                        if ch==">": depth+=1
                        else: break
                    depth=max(1,depth//2)
                    ws.row_dimensions[row].outline_level=depth
                    ws.row_dimensions[row].hidden=False
                    for col in range(1,ws.max_column+1):
                        cell=ws.cell(row=row,column=col)
                        col_name=ws.cell(row=1,column=col).value
                        if col_name not in feu_cols: cell.fill=FILL_SE
                        cell.font=FONT_SE

        # --- Impact livraison ---
        if "Impact Livraison" in headers:
            ci=headers["Impact Livraison"]
            for row in range(2,ws.max_row+1):
                v=ws.cell(row=row,column=ci).value
                if v and "RETARD" in str(v): ws.cell(row=row,column=ci).fill=FILL_ROUGE; ws.cell(row=row,column=ci).font=FONT_ROUGE
                elif v=="OK": ws.cell(row=row,column=ci).fill=FILL_VERT; ws.cell(row=row,column=ci).font=FONT_VERT
                elif v=="AUCUNE RECEPTION": ws.cell(row=row,column=ci).fill=FILL_ROUGE; ws.cell(row=row,column=ci).font=FONT_ROUGE

        # --- Hyperlinks OF -> Detail Manquants ---
        if sn=="Commandes" and "OF Associe" in headers and "Detail Manquants" in wb.sheetnames:
            co=headers["OF Associe"]; ws_det=wb["Detail Manquants"]
            hd={ws_det.cell(row=1,column=c).value:c for c in range(1,ws_det.max_column+1)}
            col_mfg=hd.get("MFGNUM")
            if col_mfg:
                mfg_rows={}
                for dr_r in range(2,ws_det.max_row+1):
                    v=ws_det.cell(row=dr_r,column=col_mfg).value
                    if v and v not in mfg_rows: mfg_rows[v]=dr_r
                for row in range(2,ws.max_row+1):
                    ov=ws.cell(row=row,column=co).value
                    if ov and str(ov).strip() in mfg_rows:
                        c=ws.cell(row=row,column=co)
                        c.hyperlink=f"#'Detail Manquants'!A{mfg_rows[str(ov).strip()]}"
                        c.font=FONT_LINK
        # --- Plan de charge: colorer les cellules taux ---
        if sn=="Plan de Charge":
            for col in range(1,ws.max_column+1):
                h=ws.cell(row=1,column=col).value
                if h and str(h).endswith("(%)"):
                    for row in range(2,ws.max_row+1):
                        v=ws.cell(row=row,column=col).value
                        if v is not None:
                            try:
                                val=float(v)
                                if val>100: style_feu(ws,row,col,"ROUGE")
                                elif val>=80: style_feu(ws,row,col,"ORANGE")
                                elif val>0: style_feu(ws,row,col,"VERT")
                            except: pass

        fmt_sheet(ws,ws.max_column)
    wb.save(f)
    print(f"\n  Export: {f}")
    return f

if __name__=="__main__":
    c,d,plan,cr,_=run()
    export_xl(c,d,plan,cr)
